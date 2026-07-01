"""Build an empty data-collection template for the DFS / shadow-economy manuscript.
Country + year are pre-filled; all indicator cells are left blank for the author to fill.
Outputs: DFS_data_template.xlsx (Codebook + Core panel + ECA panel + DiD reform dates)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

YEARS = list(range(2004, 2023))  # 2004-2022 inclusive

CORE = [
    "Kazakhstan", "Kyrgyz Republic", "Tajikistan", "Turkmenistan", "Uzbekistan",
]

# ~20 Europe & Central Asia transition economies (includes the 5 core CA states)
ECA = [
    "Albania", "Armenia", "Azerbaijan", "Belarus", "Bosnia and Herzegovina",
    "Bulgaria", "Croatia", "Georgia", "Kazakhstan", "Kyrgyz Republic",
    "Moldova", "Montenegro", "North Macedonia", "Romania", "Russian Federation",
    "Serbia", "Tajikistan", "Turkmenistan", "Ukraine", "Uzbekistan",
]

# (column key, header label, unit, maps-to-manuscript-variable)
COLS = [
    ("country",        "country",                              "name",              "identifier"),
    ("year",           "year",                                 "YYYY",              "identifier"),
    ("SHADOW",         "SHADOW",                               "% of GDP",          "Dependent var 1 (DGE estimate)"),
    ("SHADOW_MIMIC",   "SHADOW_MIMIC",                         "% of GDP",          "Alternative DV (MIMIC series) - optional"),
    ("TAX",            "TAX",                                  "% of GDP",          "Dependent var 2"),
    ("ACCOUNT",        "ACCOUNT (DFS comp i)",                 "% adults",          "DFS_INDEX component 1: account ownership"),
    ("DIGPAY",         "DIGPAY (DFS comp ii)",                 "% adults",          "DFS_INDEX component 2 AND standalone regressor"),
    ("POS_ATM",        "POS_ATM (DFS comp iii)",               "per 100k adults",   "DFS_INDEX component 3: POS+ATM density"),
    ("CASHLESS_VAL",   "CASHLESS_VAL (DFS comp iv)",           "% of GDP",          "DFS_INDEX component 4: cashless transaction value"),
    ("CASH",           "CASH",                                 "% of M2",           "Cash intensity (H3 mediator)"),
    ("GDPPC_USD",      "GDPPC_USD",                            "constant USD",      "-> I compute LGDPPC = ln(GDPPC_USD)"),
    ("OPEN",           "OPEN",                                 "% of GDP",          "Trade openness (X+M)/GDP"),
    ("UNEMP",          "UNEMP",                                "%",                 "Unemployment rate"),
    ("INFL",           "INFL",                                 "%",                 "CPI inflation"),
    ("REMIT",          "REMIT",                                "% of GDP",          "Personal remittances received"),
    ("URBAN",          "URBAN",                                "% of total",        "Urban population"),
    ("MOBILE",         "MOBILE",                               "per 100 people",    "Mobile-cellular subscriptions"),
    ("REG_QUALITY",    "REG_QUALITY",                          "WGI score (-2.5..2.5)", "-> INST component (optional if INST given)"),
    ("RULE_OF_LAW",    "RULE_OF_LAW",                          "WGI score (-2.5..2.5)", "-> INST component (optional if INST given)"),
    ("INST",           "INST",                                 "index",             "avg(REG_QUALITY, RULE_OF_LAW). Give this OR the two above"),
    ("MOBILE_COVERAGE","MOBILE_COVERAGE (IV)",                 "% pop covered",     "Instrument: 3G/4G coverage - optional"),
    ("BROADBAND",      "BROADBAND (IV)",                       "per 100 people",    "Instrument: fixed-broadband penetration - optional"),
]

hdr_fill   = PatternFill("solid", fgColor="1F4E78")
hdr_font   = Font(bold=True, color="FFFFFF", size=10)
id_fill    = PatternFill("solid", fgColor="D9E1F2")
dfs_fill   = PatternFill("solid", fgColor="FCE4D6")   # highlight DFS components
iv_fill    = PatternFill("solid", fgColor="E2EFDA")   # highlight instruments
thin       = Side(style="thin", color="BFBFBF")
border     = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "C2"  # freeze country+year and header

def build_panel(wb, title, countries):
    ws = wb.create_sheet(title)
    for j, (key, label, unit, _m) in enumerate(COLS, start=1):
        ws.cell(row=1, column=j, value=label)
        ws.column_dimensions[get_column_letter(j)].width = 20 if j > 2 else 16
    style_header(ws, len(COLS))
    r = 2
    for country in countries:
        for y in YEARS:
            ws.cell(row=r, column=1, value=country).border = border
            ws.cell(row=r, column=2, value=y).border = border
            ws.cell(row=r, column=1).fill = id_fill
            ws.cell(row=r, column=2).fill = id_fill
            for j in range(3, len(COLS) + 1):
                cell = ws.cell(row=r, column=j)
                cell.border = border
                key = COLS[j-1][0]
                if key in ("ACCOUNT", "DIGPAY", "POS_ATM", "CASHLESS_VAL"):
                    cell.fill = dfs_fill
                elif key in ("MOBILE_COVERAGE", "BROADBAND"):
                    cell.fill = iv_fill
            r += 1
    return ws

def build_codebook(wb):
    ws = wb.create_sheet("Codebook", 0)
    heads = ["Column", "Unit", "Role / maps to manuscript variable"]
    for j, h in enumerate(heads, start=1):
        ws.cell(row=1, column=j, value=h)
    style_header(ws, len(heads))
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 60
    r = 2
    for key, label, unit, mapping in COLS:
        ws.cell(row=r, column=1, value=label).border = border
        ws.cell(row=r, column=2, value=unit).border = border
        ws.cell(row=r, column=3, value=mapping).border = border
        r += 1
    # notes
    notes = [
        "",
        "HOW TO FILL:",
        "1) Leave a cell blank if you do not have that value (do NOT put 0). Blank = missing.",
        "2) Orange columns = the 4 raw DFS components. If you give these, I compute the PCA (Table 4: KMO, Bartlett, eigenvalue, loadings) and DFS_INDEX myself.",
        "   If you already have a finished DFS_INDEX instead, add a column 'DFS_INDEX' and I will use it, but then I cannot reproduce PCA diagnostics.",
        "3) Green columns = instruments for the IV/2SLS (optional but needed for Table 6 IV row).",
        "4) You may give INST directly, OR give REG_QUALITY and RULE_OF_LAW and I will average them.",
        "5) GDPPC_USD in constant USD; I take the natural log for LGDPPC.",
        "6) Core panel = 5 Central Asian states. ECA panel = ~20 transition economies (superset).",
        "   You can fill only the Core sheet if the ECA data is hard to assemble.",
    ]
    for line in notes:
        ws.cell(row=r, column=1, value=line)
        if line.endswith(":"):
            ws.cell(row=r, column=1).font = Font(bold=True)
        r += 1

def build_did(wb):
    ws = wb.create_sheet("DiD_reform_dates")
    heads = ["country", "reform_name", "reform_year", "notes"]
    for j, h in enumerate(heads, start=1):
        ws.cell(row=1, column=j, value=h)
    style_header(ws, len(heads))
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 40
    r = 2
    for country in CORE:
        ws.cell(row=r, column=1, value=country).border = border
        for j in range(2, 5):
            ws.cell(row=r, column=j).border = border
        r += 1
    ws.cell(row=r+1, column=1,
            value="Enter the year of the main payment-system / financial reform per country "
                  "(e.g. Kazakhstan super-app; Uzbekistan 2017 liberalisation). Used for the DiD design.")

wb = openpyxl.Workbook()
wb.remove(wb.active)  # drop default sheet
build_codebook(wb)
build_panel(wb, "Core_panel_CA", CORE)
build_panel(wb, "ECA_panel", ECA)
build_did(wb)
wb.save("DFS_data_template.xlsx")
print("WROTE DFS_data_template.xlsx")
print("Core rows:", len(CORE)*len(YEARS), "| ECA rows:", len(ECA)*len(YEARS), "| columns:", len(COLS))
