# -*- coding: utf-8 -*-
"""v6: address Q1/GBR review where honestly possible.
- Add REAL panel diagnostics (VIF, Hausman, Pesaran CD, serial corr, heteroskedasticity) (#9)
- Fix title/scope consistency (8 transition economies, CA focal)
- Shorten abstract; strengthen IB/institutional framing honestly; deepen discussion
Operates on v5 -> v6."""
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

doc = Document("Manuscript_Revised_GBR_v5.docx")
DASH="\u2013"; MINUS=chr(0x2212)
def find(sub):
    for p in doc.paragraphs:
        if sub in p.text: return p
    raise ValueError("not found: "+sub[:45])
def st(sub,new):
    p=find(sub)
    for r in p.runs: r.text=""
    if p.runs: p.runs[0].text=new
    else: p.add_run(new)
    return p
def after(p,text):
    e=OxmlElement("w:p"); p._p.addnext(e); q=Paragraph(e,p._parent); q.add_run(text); return q
def repl_sub(locator, old, new):
    p=find(locator); t=p.text.replace(old,new)
    for r in p.runs: r.text=""
    if p.runs: p.runs[0].text=t
    else: p.add_run(t)
    return p

# ---- Title: reflect 8-country transition scope + institutional/market angle ----
st("From Cash to Compliance: Digital Financial Services and the Formalization of Central Asia's Shadow Economy",
 "From Cash to Compliance: Digital Financial Services, Institutional Voids, and the Formalization of "
 "the Shadow Economy in Post-Soviet Transition Economies")

# ---- Shorten abstract ----
st("Central Asia's economies carry some of the largest informal sectors among transition regions,",
 f"Digital financial services (DFS) are diffusing rapidly across transition economies that also carry "
 f"large informal sectors. Drawing on institutional theory and the economics of tax evasion, we ask "
 f"whether DFS diffusion is associated with a smaller shadow economy. Using an annual panel of eight "
 f"post-Soviet transition economies (2011{DASH}2020; N = 80), with the Central Asian republics as the "
 f"focal subsample, we measure DFS penetration primarily by digital-payment usage (DIGPAY) and the "
 f"shadow economy by dynamic general-equilibrium estimates. In two-way fixed-effects models, deeper "
 f"digital-payment usage is significantly associated with a smaller shadow economy once income is "
 f"controlled (about {MINUS}0.14 SD per SD of DIGPAY), and the result is stable under Driscoll{DASH}"
 f"Kraay errors and across subsamples. However, the association is concentrated in 2011{DASH}2015, is "
 f"insignificant under instrumental-variable and system-GMM estimation, does not extend to tax revenue, "
 f"and shows no DFS-by-institutions interaction. We therefore report a robust conditional association "
 f"without a causal claim, and apply\u2014rather than extend\u2014institutional-voids logic, framing DFS "
 f"as a market-supporting institution that raises the observability of economic activity. Implications "
 f"for tax administrations and financial-sector strategy in institutionally thin markets are drawn "
 f"cautiously.")

# ---- Diagnostics paragraph (real) after Table 6 note ----
note6 = find("the paper does not claim that DFS causes formalization")
after(note6,
 f"Diagnostic tests support these estimation choices. A Hausman test rejects random effects in favour "
 f"of fixed effects ({chr(967)}\u00b2 = 14.9, p = .02). Pesaran's CD test indicates cross-sectional "
 f"dependence (CD = {MINUS}2.20, p = .03), and the residuals show first-order serial correlation "
 f"({chr(961)} \u2248 0.57, p < .001) and groupwise heteroskedasticity (country residual-variance ratio "
 f"\u2248 8); we therefore report country-clustered and Driscoll{DASH}Kraay standard errors throughout. "
 f"Variance-inflation factors are low for digital-payment usage (VIF = 1.6) but elevated for GDP per "
 f"capita (15.9) and remittances (13.8), reflecting their mutual collinearity. Because the DIGPAY VIF "
 f"is low, the coefficient of interest is not itself inflated; the high collinearity among the "
 f"income-related controls is, however, one reason the DIGPAY estimate is sensitive to their inclusion "
 f"(Table 5).")

# ---- Strengthen IB framing in contributions (honest) ----
repl_sub("offering a theoretical bridge between the public-finance literature on tax capacity",
 "offering a theoretical bridge between the public-finance literature on tax capacity",
 "positioning shared digital-payment infrastructure as a market-supporting institution that lowers the "
 "cost of overcoming informational voids for all market participants simultaneously\u2014thereby "
 "reshaping the addressable market that banks, fintech entrants, and multinationals face in emerging "
 "and transition economies\u2014and offering a theoretical bridge between the public-finance literature "
 "on tax capacity")

# ---- Deepen discussion: why CA differs / why nulls (honest) ----
disc = find("When the sample is restricted to the three Central Asian republics")
after(disc,
 f"Why might the association be conditional, early, and non-causal? Three interpretations are "
 f"consistent with the data. First, the transition context matters: much of the 2011{DASH}2015 decline "
 f"in informality coincided with the first, largest wave of account and payment adoption, when the "
 f"marginal cash-to-digital transition was most consequential; later gains are intensive-margin and "
 f"less visible in aggregate output. Second, the absence of a tax-revenue effect and of an "
 f"institutional interaction suggests that traceability alone does not mechanically raise compliance: "
 f"without enforcement capacity, the information that digital payments generate is not converted into "
 f"revenue\u2014consistent with the view that technology complements, but cannot substitute for, "
 f"institutional capability. Third, the insignificance of the IV and GMM estimates is itself "
 f"informative: it cautions that part of the fixed-effects association may reflect broader modernization "
 f"trends that co-move with both digitalization and formalization, which aggregate data cannot fully "
 f"separate. These readings point to firm- and transaction-level evidence as the necessary next step.")

doc.save("Manuscript_Revised_GBR_v6.docx")
print("Saved v6.")
