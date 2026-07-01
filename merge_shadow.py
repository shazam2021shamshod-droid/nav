"""Merge SHADOW (DGE) and SHADOW_MIMIC from the World Bank Informal Economy Database
into DFS_data_FILLED.xlsx. Only real values; missing years left blank.
"""
import openpyxl
from openpyxl.styles import PatternFill

YEARS = list(range(2004, 2023))
ISO_NEEDED = {"KAZ","KGZ","TJK","TKM","UZB","ALB","ARM","AZE","BLR","BIH","BGR",
              "HRV","GEO","MDA","MNE","MKD","ROU","RUS","SRB","UKR"}

def read_sheet(wb, sheet):
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    # map column index -> year (int) where header is a 4-digit year
    ycol = {}
    for j, h in enumerate(header):
        s = str(h).strip() if h is not None else ""
        if s.isdigit() and len(s) == 4:
            ycol[j] = int(s)
    out = {}  # iso -> {year: value}
    for r in rows[1:]:
        if not r or len(r) < 2:
            continue
        code = (str(r[1]).strip() if r[1] else "")
        if code not in ISO_NEEDED:
            continue
        d = {}
        for j, yr in ycol.items():
            if yr in YEARS and j < len(r):
                v = r[j]
                if v not in (None, ""):
                    try:
                        d[yr] = round(float(v), 4)
                    except (TypeError, ValueError):
                        pass
        out[code] = d
    return out

src = openpyxl.load_workbook("informal-economy-database.xlsx", read_only=True, data_only=True)
dge = read_sheet(src, "DGE_p")
mimic = read_sheet(src, "MIMIC_p")
print("DGE countries:", len(dge), "| MIMIC countries:", len(mimic))
print("DGE last year available (KAZ):", max(dge["KAZ"]) if dge.get("KAZ") else None)

# name -> iso for the FILLED workbook rows
NAME2ISO = {
    "Kazakhstan":"KAZ","Kyrgyz Republic":"KGZ","Tajikistan":"TJK","Turkmenistan":"TKM",
    "Uzbekistan":"UZB","Albania":"ALB","Armenia":"ARM","Azerbaijan":"AZE","Belarus":"BLR",
    "Bosnia and Herzegovina":"BIH","Bulgaria":"BGR","Croatia":"HRV","Georgia":"GEO",
    "Moldova":"MDA","Montenegro":"MNE","North Macedonia":"MKD","Romania":"ROU",
    "Russian Federation":"RUS","Serbia":"SRB","Ukraine":"UKR","Uzbekistan ":"UZB",
}

green = PatternFill("solid", fgColor="C6EFCE")
wb = openpyxl.load_workbook("DFS_data_FILLED.xlsx")
filled_dge = filled_mimic = 0
for sh in ["Core_panel_CA", "ECA_panel"]:
    ws = wb[sh]
    header = [c.value for c in ws[1]]
    # locate SHADOW / SHADOW_MIMIC columns by prefix
    def colidx(prefix):
        for j, h in enumerate(header):
            if h and str(h).startswith(prefix):
                return j
        return None
    c_country = 0
    c_shadow = colidx("SHADOW (") if colidx("SHADOW (") is not None else colidx("SHADOW")
    c_mimic = colidx("SHADOW_MIMIC")
    for row in ws.iter_rows(min_row=2):
        name = row[c_country].value
        yr = row[1].value
        iso = NAME2ISO.get(str(name).strip())
        if not iso:
            continue
        if c_shadow is not None and iso in dge and yr in dge[iso]:
            cell = row[c_shadow]; cell.value = dge[iso][yr]; cell.fill = green; filled_dge += 1
        if c_mimic is not None and iso in mimic and yr in mimic[iso]:
            cell = row[c_mimic]; cell.value = mimic[iso][yr]; cell.fill = green; filled_mimic += 1

wb.save("DFS_data_FILLED.xlsx")
print(f"Filled SHADOW (DGE): {filled_dge}  |  SHADOW_MIMIC: {filled_mimic}")
